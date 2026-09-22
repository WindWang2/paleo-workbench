"""Tabular interchange adapters: CSV / TSV / Excel.

Adds what the preview parsers never provided: delimiter/encoding sniffing,
header/column diagnostics (duplicates, mixed types, coordinate/depth/well-name
aliases), a serializable schema-mapping *preset* (no absolute paths), and a
normalizing transform import (delimiter/encoding canonicalization) whose
output is verified before it reaches the catalog.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from paleo_workbench.interchange.contracts import (
    CancelToken,
    ExportPlan,
    ExportVerification,
    FormatCapability,
    ImportExecutionResult,
    ImportPlan,
    InspectionResult,
    NULL_CANCEL,
    VerificationCheck,
    VerificationState,
)
from paleo_workbench.interchange.registry import FormatAdapter, SNIFF_PREFIX_BYTES

# Bounded analysis budget for sniffing/counting (never load whole files).
_SAMPLE_BYTES = 256 * 1024
_COUNT_ROWS_MAX = 2_000_000

_DELIMITERS = (",", "\t", ";", "|")

CANONICAL_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "well_name": ("well", "wellname", "井名", "井号", "well_id", "borehole"),
    "depth": ("depth", "md", "深度", "井深", "measured_depth"),
    "x": ("x", "east", "easting", "x坐标", "经度", "lng", "lon", "long", "longitude"),
    "y": ("y", "north", "northing", "y坐标", "纬度", "lat", "latitude"),
    "value": ("value", "z", "值", "数值", "factor", "属性"),
}


@dataclass
class TabularMappingPreset:
    """Serializable column-mapping preset. Contains NO absolute paths."""

    delimiter: str = ","
    encoding: str = "utf-8"
    has_header: bool = True
    header_map: dict[str, str] = field(default_factory=dict)  # canonical -> source column
    decimal: str = "."
    units: dict[str, str] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "delimiter": self.delimiter,
            "encoding": self.encoding,
            "has_header": self.has_header,
            "header_map": dict(self.header_map),
            "decimal": self.decimal,
            "units": dict(self.units),
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "TabularMappingPreset":
        return cls(
            delimiter=str(data.get("delimiter", ",")),
            encoding=str(data.get("encoding", "utf-8")),
            has_header=bool(data.get("has_header", True)),
            header_map={str(k): str(v) for k, v in data.get("header_map", {}).items()},
            decimal=str(data.get("decimal", ".")),
            units={str(k): str(v) for k, v in data.get("units", {}).items()},
        )


def sniff_delimiter(sample: str) -> str:
    """Pick the delimiter that yields the most consistent column count."""
    lines = [ln for ln in sample.splitlines() if ln.strip()][:32]
    if not lines:
        return ","
    best, best_score = ",", -1.0
    for delim in _DELIMITERS:
        try:
            rows = list(csv.reader(lines, delimiter=delim))
        except Exception:
            continue
        if len(rows) < 2:
            continue
        counts = [len(r) for r in rows]
        if max(counts) < 2:
            continue
        mean = sum(counts) / len(counts)
        variance = sum((c - mean) ** 2 for c in counts) / len(counts)
        score = (mean - 1) / (1.0 + variance)
        if score > best_score:
            best, best_score = delim, score
    return best


def _canonical_column(name: str) -> str | None:
    key = name.strip().casefold()
    for canonical, aliases in CANONICAL_COLUMN_ALIASES.items():
        if key in aliases:
            return canonical
    return None


class CsvLikeAdapter(FormatAdapter):
    resource_type = "tabular"
    default_delimiter = ","

    def capability(self) -> FormatCapability:
        return FormatCapability(
            read=True,
            inspect=True,
            import_data=True,
            export=True,
            roundtrip_verify=True,
            notes="支持 delimiter/encoding 探测与归一化转换导入；导出为归一化 CSV",
        )

    # -- helpers ------------------------------------------------------------
    def _decode_sample(self, path: Path, result: InspectionResult) -> tuple[str, str]:
        """Return (text_sample, encoding) with a GB18030 fallback like the
        existing preview parsers (UTF-8-sig -> GB18030 -> replace)."""
        raw = b""
        try:
            with open(path, "rb") as fh:
                raw = fh.read(_SAMPLE_BYTES)
        except OSError as exc:
            result.ok = False
            result.errors.append(f"无法读取文件: {exc}")
            return "", "utf-8"
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return raw.decode(encoding), encoding
            except UnicodeDecodeError:
                continue
        result.warnings.append("文件编码无法可靠识别（已按替换模式解码）")
        return raw.decode("utf-8", errors="replace"), "utf-8"

    def _sniff(self, path: Path) -> str:
        try:
            with open(path, "rb") as fh:
                raw = fh.read(_SAMPLE_BYTES)
        except OSError:
            return self.default_delimiter
        text = raw.decode("utf-8", errors="replace")
        return sniff_delimiter(text)

    # -- inspect ------------------------------------------------------------
    def inspect(self, path: Path) -> InspectionResult:
        path = Path(path)
        result = InspectionResult(format_id=self.format_id, object_type=self.resource_type)
        try:
            result.size_bytes = path.stat().st_size
        except OSError as exc:
            result.ok = False
            result.errors.append(f"无法读取文件状态: {exc}")
            return result

        sample, encoding = self._decode_sample(path, result)
        if not result.ok:
            return result
        if encoding not in ("utf-8", "utf-8-sig"):
            result.warnings.append(f"检测到非 UTF-8 编码: {encoding}")
        delimiter = sniff_delimiter(sample)
        rows = list(csv.reader(sample.splitlines(), delimiter=delimiter))
        rows = [r for r in rows if any(cell.strip() for cell in r)]
        if not rows:
            result.warnings.append("未检测到数据行")
            result.metadata = {"delimiter": delimiter, "encoding": encoding, "columns": []}
            return result

        columns = [c.strip() for c in rows[0]]
        dupes = sorted({c for c in columns if c and columns.count(c) > 1})
        if dupes:
            result.warnings.append(f"重复表头列名: {', '.join(dupes)}")

        has_header = self._looks_like_header(rows[0], rows[1:])
        body = rows[1:] if has_header else rows
        if rows[1:] and not has_header:
            result.warnings.append("首行不像表头（可能缺少表头）")

        widths = {len(r) for r in rows}
        if len(widths) > 1:
            result.warnings.append(f"列数不一致（{min(widths)}~{max(widths)} 列）")

        decimal = "."
        sample_cells = [cell for row in body[:64] for cell in row if cell.strip()]
        if sample_cells:
            numeric_dot = sum(1 for c in sample_cells if self._is_numeric(c, "."))
            numeric_comma = sum(1 for c in sample_cells if self._is_numeric(c, ","))
            if numeric_comma > numeric_dot and numeric_comma / len(sample_cells) >= 0.5:
                decimal = ","
                result.warnings.append("检测到逗号小数（decimal=','）")
            elif numeric_dot / len(sample_cells) < 0.5:
                result.warnings.append("数据主体非数值（表格可能为文本/混合类型）")

        header_map: dict[str, str] = {}
        for column in columns:
            canonical = _canonical_column(column)
            if canonical and canonical not in header_map:
                header_map[canonical] = column
        if "well_name" not in header_map and columns:
            result.warnings.append("未识别到井名列（well/井名 别名表未命中）")
        if ("x" in header_map) != ("y" in header_map):
            result.warnings.append("坐标列不完整（仅识别到 X 或 Y 之一）")

        row_count = self._count_data_rows(path, delimiter, has_header=has_header)
        result.metadata = {
            "delimiter": delimiter,
            "encoding": encoding,
            "columns": columns,
            "header_map": header_map,
            "has_header": has_header,
            "row_count": row_count,
            "decimal": decimal,
        }
        return result

    @staticmethod
    def _strings_only(row: list[str]) -> bool:
        non_empty = [cell for cell in row if cell.strip()]
        return bool(non_empty) and all(_cell_is_text(cell) for cell in non_empty)

    @staticmethod
    def _looks_like_header(header: list[str], body: list[list[str]]) -> bool:
        if not body:
            return True
        return CsvLikeAdapter._strings_only(header) and not CsvLikeAdapter._strings_only(body[0])

    @staticmethod
    def _is_numeric(cell: str, decimal: str) -> bool:
        text = cell.strip()
        if not text:
            return False
        probe = text.replace(decimal, ".", 1) if decimal != "." else text
        try:
            float(probe)
            return True
        except ValueError:
            return False

    @staticmethod
    def _count_data_rows(path: Path, delimiter: str, *, has_header: bool = True) -> int:
        count = 0
        try:
            with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
                reader = csv.reader(fh, delimiter=delimiter)
                for i, row in enumerate(reader):
                    if i >= _COUNT_ROWS_MAX:
                        break
                    if any(cell.strip() for cell in row):
                        count += 1
        except OSError:
            return count
        return max(0, count - (1 if has_header else 0))

    # -- plan / execute -----------------------------------------------------
    def plan_import(self, path, inspection, *, managed=True, asset_name=None, options=None):
        plan = super().plan_import(
            path, inspection, managed=managed, asset_name=asset_name, options=options
        )
        preset = TabularMappingPreset(
            delimiter=inspection.metadata.get("delimiter", self.default_delimiter),
            encoding=inspection.metadata.get("encoding", "utf-8"),
            has_header=bool(inspection.metadata.get("has_header", True)),
            header_map=dict(inspection.metadata.get("header_map", {})),
            decimal=inspection.metadata.get("decimal", "."),
        )
        plan.metadata["mapping_preset"] = preset.to_dict()
        # A preset travels with the plan but never carries absolute paths.
        plan.options["mapping_preset"] = preset.to_dict()
        return plan

    def import_data(self, path, plan, *, work_dir, catalog, cancel=None, progress=None) -> ImportExecutionResult:
        cancel = cancel or NULL_CANCEL
        cancel.checkpoint()
        preset = TabularMappingPreset.from_dict(plan.options.get("mapping_preset", {}))
        needs_transform = (
            preset.delimiter != ","
            or preset.encoding not in ("utf-8", "utf-8-sig")
            or plan.options.get("normalize")
        )
        source = Path(plan.source_path)
        if not needs_transform:
            version = catalog.import_raw(
                source,
                name=plan.asset_name,
                type=self.resource_type,
                format=self.format_id,
                metadata=plan.metadata,
            )
            return ImportExecutionResult(version_id=version.id, asset_id=version.asset_id, managed=True)

        # Normalize into the work dir, verify, then register the normalized copy.
        work_dir = Path(work_dir)
        work_dir.mkdir(parents=True, exist_ok=True)
        staged = work_dir / f"{source.stem}.normalized.csv"
        try:
            self._normalize(source, staged, preset)
            verification = self._verify_normalized_csv(staged, preset)
            if not verification.ok:
                raise RuntimeError(f"归一化输出未通过校验: {verification.detail}")
            cancel.checkpoint()
            version = catalog.import_raw(
                staged,
                name=plan.asset_name,
                type=self.resource_type,
                format="csv",
                metadata={**plan.metadata, "normalized_from": self.format_id},
            )
        finally:
            try:
                staged.unlink(missing_ok=True)
            except OSError:
                pass
        return ImportExecutionResult(
            version_id=version.id, asset_id=version.asset_id, managed=True,
            verification=verification,
        )

    def _normalize(self, source: Path, target: Path, preset: TabularMappingPreset) -> None:
        with open(source, "r", encoding=preset.encoding, errors="replace", newline="") as src:
            reader = csv.reader(src, delimiter=preset.delimiter)
            rows = (row for row in reader if any(cell.strip() for cell in row))
            with open(target, "w", encoding="utf-8", newline="") as dst:
                writer = csv.writer(dst)
                for row in rows:
                    writer.writerow(row)

    def _verify_normalized_csv(self, target: Path, preset: TabularMappingPreset) -> ExportVerification:
        checks: list[VerificationCheck] = []
        if not target.is_file() or target.stat().st_size == 0:
            return ExportVerification(VerificationState.FAILED, checks, detail="归一化输出为空")
        with open(target, "r", encoding="utf-8", newline="") as fh:
            rows = list(csv.reader(fh))
        checks.append(VerificationCheck("utf8_parsable", True))
        if not rows:
            return ExportVerification(VerificationState.FAILED, checks, detail="无数据行")
        widths = {len(r) for r in rows}
        checks.append(VerificationCheck("consistent_columns", len(widths) == 1, f"列宽 {sorted(widths)}"))
        with open(target, "rb") as fh:
            head = fh.read(3)
        checks.append(VerificationCheck("no_bom", not head.startswith(b"\xef\xbb\xbf")))
        failed = [c for c in checks if not c.passed]
        if failed:
            return ExportVerification(VerificationState.FAILED, checks)
        return ExportVerification(VerificationState.VERIFIED, checks)

    # -- export: normalized CSV --------------------------------------------
    def plan_export(self, source_path, target_path, *, options=None):
        target_path = Path(target_path)
        if target_path.suffix.lower() not in (".csv", ".tsv"):
            from paleo_workbench.interchange.contracts import FormatNotSupportedError

            raise FormatNotSupportedError("表格适配器仅支持导出为 CSV/TSV")
        return ExportPlan(
            format_id=self.format_id,
            source_path=str(source_path),
            target_path=str(target_path),
            estimated_bytes=Path(source_path).stat().st_size,
            options=dict(options or {}),
        )

    def export_data(self, source_path, plan, *, work_dir, cancel=None, progress=None):
        cancel = cancel or NULL_CANCEL
        cancel.checkpoint()
        from paleo_workbench.resources.exporters import atomic_output

        source = Path(source_path)
        target = Path(plan.target_path)
        # Read with the SOURCE's real delimiter (sniffed — the extension does
        # not decide, a .csv can be semicolon-delimited), write with the
        # target's convention.
        source_delimiter = sniff_delimiter(self._read_text_prefix(source))
        target_delimiter = "\t" if target.suffix.lower() == ".tsv" else ","
        with atomic_output(target) as tmp:
            with open(source, "r", encoding="utf-8", errors="replace", newline="") as src:
                reader = csv.reader(src, delimiter=source_delimiter)
                with open(tmp, "w", encoding="utf-8", newline="") as dst:
                    writer = csv.writer(dst, delimiter=target_delimiter)
                    for row in reader:
                        writer.writerow(row)
        return target

    @staticmethod
    def _read_text_prefix(path: Path, limit: int = _SAMPLE_BYTES) -> str:
        with open(path, "rb") as fh:
            return fh.read(limit).decode("utf-8", errors="replace")

    def verify_output(self, target_path, plan) -> ExportVerification:
        target_path = Path(target_path)
        checks: list[VerificationCheck] = []
        if not target_path.is_file():
            return ExportVerification(VerificationState.FAILED, checks, detail="输出文件不存在")
        if target_path.stat().st_size == 0:
            return ExportVerification(VerificationState.FAILED, checks, detail="输出文件为空")
        delimiter = "\t" if target_path.suffix.lower() == ".tsv" else ","
        with open(target_path, "r", encoding="utf-8", newline="") as fh:
            rows = list(csv.reader(fh, delimiter=delimiter))
        checks.append(VerificationCheck("parsable", bool(rows), f"{len(rows)} 行"))
        widths = {len(r) for r in rows}
        checks.append(VerificationCheck("consistent_columns", len(widths) == 1))
        src_delim = "," if delimiter == "\t" else "\t"
        source_rows = _count_rows_fast(Path(plan.source_path), src_delim)
        checks.append(
            VerificationCheck(
                "row_count_within_1",
                abs(len(rows) - source_rows) <= 1,
                f"源约 {source_rows} 行，输出 {len(rows)} 行",
            )
        )
        failed = [c for c in checks if not c.passed]
        if failed:
            return ExportVerification(VerificationState.FAILED, checks)
        return ExportVerification(VerificationState.VERIFIED, checks)


def _count_rows_fast(path: Path, delimiter: str) -> int:
    try:
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
            return sum(1 for row in csv.reader(fh, delimiter=delimiter) if any(cell.strip() for cell in row))
    except OSError:
        return 0


def _cell_is_text(cell: str) -> bool:
    try:
        float(cell.strip())
        return False
    except (ValueError, TypeError):
        return True


class CsvAdapter(CsvLikeAdapter):
    format_id = "csv"
    display_name = "CSV 表格"
    extensions = ("csv",)
    default_delimiter = ","


class TsvAdapter(CsvLikeAdapter):
    format_id = "tsv"
    display_name = "TSV 表格"
    extensions = ("tsv", "txt")
    default_delimiter = "\t"


class ExcelAdapter(FormatAdapter):
    format_id = "xlsx"
    display_name = "Excel 工作簿"
    extensions = ("xlsx", "xls")
    resource_type = "spreadsheet"

    def capability(self) -> FormatCapability:
        return FormatCapability(
            read=True,
            inspect=True,
            import_data=True,
            export=False,
            roundtrip_verify=False,
            notes="按工作簿整体导入（保留原文件）；导出走既有 table_to_xlsx 转换链",
        )

    def inspect(self, path: Path) -> InspectionResult:
        path = Path(path)
        result = InspectionResult(format_id=self.format_id, object_type=self.resource_type)
        try:
            result.size_bytes = path.stat().st_size
        except OSError as exc:
            result.ok = False
            result.errors.append(f"无法读取文件状态: {exc}")
            return result
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            result.ok = False
            result.errors.append(f"openpyxl 不可用: {exc}")
            return result
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            result.ok = False
            result.errors.append(f"工作簿打开失败（文件可能损坏）: {exc}")
            return result
        sheets: list[dict[str, Any]] = []
        try:
            for name in workbook.sheetnames:
                sheet = workbook[name]
                info: dict[str, Any] = {
                    "name": name,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                }
                if sheet.max_row and sheet.max_row > 0:
                    first_row = next(
                        sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()
                    )
                    columns = [str(c).strip() if c is not None else "" for c in first_row]
                    info["columns"] = columns
                    dupes = sorted({c for c in columns if c and columns.count(c) > 1})
                    if dupes:
                        result.warnings.append(f"工作表 {name}: 重复列名 {', '.join(dupes)}")
                    header_map = {}
                    for column in columns:
                        canonical = _canonical_column(column)
                        if canonical and canonical not in header_map:
                            header_map[canonical] = column
                    info["header_map"] = header_map
                sheets.append(info)
        finally:
            workbook.close()
        result.metadata = {"sheets": sheets}
        if len(sheets) > 1:
            result.warnings.append(f"多工作表（{len(sheets)}）：导入保留整个工作簿")
        return result

    def plan_import(self, path, inspection, *, managed=True, asset_name=None, options=None):
        plan = super().plan_import(
            path, inspection, managed=managed, asset_name=asset_name, options=options
        )
        plan.metadata.setdefault("object_type", self.resource_type)
        return plan

    def import_data(self, path, plan, *, work_dir, catalog, cancel=None, progress=None) -> ImportExecutionResult:
        cancel = cancel or NULL_CANCEL
        cancel.checkpoint()
        version = catalog.import_raw(
            Path(path),
            name=plan.asset_name,
            type=self.resource_type,
            format=self.format_id,
            metadata=plan.metadata,
        )
        return ImportExecutionResult(version_id=version.id, asset_id=version.asset_id, managed=True)
