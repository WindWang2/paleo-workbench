"""Freeze resources/export_service.py + data_asset_registry.py behavior.

Drives the real Python implementation over a scripted temp tree and freezes
the observable outputs: ExportJobResult fields, appended export_artifacts
rows, written file bytes (as text), classify/scan/export registry results.

Non-deterministic fields are normalized before freezing:
  * absolute paths → fixture-relative (``<root>``-anchored)
  * artifact ids ("artifact_<12hex>") → "artifact_<id>"
  * generated_at timestamps → "<ts>"
  * catalog registration: PALEO_DATA_CATALOG is unset in the harness →
    get_catalog() is None → catalog_version_id stays None (the injected-
    registrar branch is exercised on the C++ side with a fake registrar).
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

import _legacy_reference

_legacy_reference.ensure_legacy_reference()  # archived-reference shim

from paleo_workbench.project.models import (  # noqa: E402
    ExportArtifact,
    ProjectDocument,
    ProjectMeta,
    ResourceItem,
)
from paleo_workbench.resources.export_service import (  # noqa: E402
    default_export_dir,
    export_asset_to_path,
    export_project_inventory,
    list_asset_export_labels,
    register_exported_view,
    _view_format_rank,
)
from paleo_workbench.resources.data_asset_registry import (  # noqa: E402
    DataAssetRegistry,
    FormatSpec,
)

OUT = (
    Path(__file__).resolve().parents[2]
    / "libs"
    / "ui_data_core"
    / "ui_data_core_tests"
    / "fixtures"
    / "export_oracle.json"
)


def _norm_path(p: str, root: Path) -> str:
    pp = Path(p)
    if pp.is_absolute():
        try:
            return pp.resolve().relative_to(root.resolve()).as_posix()
        except (ValueError, OSError):
            return pp.name
    return p


def _norm_artifact(a) -> dict | None:
    if a is None:
        return None
    d = a.model_dump()
    d["id"] = "artifact_<id>"
    d["generated_at"] = "<ts>"
    d["output_path"] = _norm_path(d["output_path"], _ROOT)
    return d


def _norm_result(r) -> dict:
    return {
        "success": r.success,
        "output_path": _norm_path(r.output_path, _ROOT),
        "format": r.format,
        "artifact": _norm_artifact(r.artifact),
        "message": _norm_message(r.message),
        "warnings": list(r.warnings),
    }


def _norm_message(m: str, ) -> str:
    # "<abs path>" inside a message → relative
    out = m
    text = str(_ROOT.resolve())
    if text in out:
        out = out.replace(text, "<root>")
    return out


def _project_artifacts(project: ProjectDocument) -> list[dict]:
    return [_norm_artifact(a) for a in project.export_artifacts]


_ROOT: Path


def main() -> None:
    global _ROOT
    _ROOT = Path(tempfile.mkdtemp(prefix="export_oracle_"))
    tree = _ROOT / "tree"
    tree.mkdir()

    files = {
        "table.csv": "a,b\n1,2\n3,4\n",
        "notes.txt": "line1\nline2\n",
        "bad.las": "not a las\n",
        "map.geojson": '{"type":"FeatureCollection","features":[]}',
        "x.xyz": "unknown\n",
    }
    for rel, content in files.items():
        (tree / rel).write_text(content, encoding="utf-8")

    cases = []

    # --- list_asset_export_labels / view_format_rank ---------------------
    labels = {}
    for fmt in ("csv", "las", "geojson", "png", "txt", "xyz", "sgy"):
        labels[fmt] = list_asset_export_labels(
            ResourceItem(name="x." + fmt, path=str(tree / ("x." + fmt)),
                         type="tabular", format=fmt)
        )
    ranks = {lbl: _view_format_rank(lbl) for lbl in
             ("PNG", "svg", "Pdf", "JSON", "", "png")}
    cases.append({"id": "labels", "labels": labels, "ranks": ranks})

    # --- export_asset_to_path -------------------------------------------
    def asset(name, fmt, rtype="tabular"):
        return ResourceItem(id="res_src01", name=name,
                            path=str(tree / name), type=rtype, format=fmt)

    def run_export(case_id, asset_, label, out_name, project=None,
                   project_path=None):
        out = tree / out_name
        r = export_asset_to_path(
            asset_, label, out, project=project, project_path=project_path
        )
        entry = {"id": case_id, "result": _norm_result(r)}
        if out.exists():
            if out.suffix == ".xlsx":
                import openpyxl

                wb = openpyxl.load_workbook(out)
                ws = wb[wb.sheetnames[0]]
                entry["output_xlsx"] = {
                    "sheet": wb.sheetnames[0],
                    "grid": [
                        [c.value for c in row] for row in ws.iter_rows()
                    ],
                }
            else:
                entry["output_text"] = out.read_text(encoding="utf-8")
        return entry

    csv_asset = asset("table.csv", "csv")
    cases.append(run_export("asset.csv_to_json", csv_asset, "JSON",
                            "out_table.json"))
    cases.append(run_export("asset.csv_to_xlsx", csv_asset, "XLSX",
                            "out_table.xlsx"))
    cases.append(run_export("asset.unsupported_label", csv_asset, "PDF",
                            "out.pdf"))
    missing = asset("missing.csv", "csv")
    cases.append(run_export("asset.missing_source", missing, "JSON",
                            "out_missing.json"))
    cases.append(run_export("asset.las_error", asset("bad.las", "las",
                                                   "well_log"), "CSV",
                            "out_bad.csv"))

    # with project + project_path → artifact registered, path relativized
    proj_dir = _ROOT / "proj"
    proj_dir.mkdir()
    project_path = proj_dir / "project.pwb"
    project_path.write_text("{}", encoding="utf-8")
    project = ProjectDocument(meta=ProjectMeta(name="demo", region="R1"))
    r = export_asset_to_path(csv_asset, "JSON", tree / "out_reg.json",
                             project=project, project_path=project_path)
    cases.append({
        "id": "asset.registered",
        "result": _norm_result(r),
        "output_text": (tree / "out_reg.json").read_text(encoding="utf-8"),
        "project_artifacts": _project_artifacts(project),
    })

    # --- export_project_inventory ---------------------------------------
    project2 = ProjectDocument(meta=ProjectMeta(name="inv", region="Z"))
    project2.resources.append(csv_asset)
    project2.resources.append(
        ResourceItem(id="res_gj", name="map.geojson",
                     path=str(tree / "map.geojson"), type="geojson",
                     format="geojson", artifact_role="input",
                     tags=["input"], parsed_summary={"k": 1}))
    project2.export_artifacts.append(
        ExportArtifact(id="artifact_fixed01", linked_id="res_src01",
                       format="json", output_path="out_table.json"))
    inv_out = tree / "inv.json"
    seed_project = project2.model_dump(mode="json")
    for row in seed_project["resources"]:
        row["path"] = _norm_path(row["path"], _ROOT)
    r = export_project_inventory(project2, inv_out, project_path=project_path)
    written = json.loads(inv_out.read_text(encoding="utf-8"))
    for row in written["resources"]:
        row["path"] = _norm_path(row["path"], _ROOT)
    cases.append({
        "id": "inventory",
        "result": _norm_result(r),
        "output_json": written,
        "project_artifacts": _project_artifacts(project2),
        "seed_project": seed_project,
    })

    # --- register_exported_view ------------------------------------------
    project3 = ProjectDocument(meta=ProjectMeta(name="viz", region=""))
    art = register_exported_view(
        object(), tree / "view.png", "PNG", project=project3,
        project_path=project_path, linked_id="viz_view",
        source_task_ids=["task_a", "task_b"])
    cases.append({
        "id": "register_view",
        "artifact": _norm_artifact(art),
        "project_artifacts": _project_artifacts(project3),
    })
    # register=False → None
    assert register_exported_view(object(), tree / "v.png", "PNG",
                                  project=project3, register=False) is None

    # --- default_export_dir ----------------------------------------------
    cases.append({
        "id": "default_export_dir",
        "no_project": default_export_dir(None).as_posix().rsplit("/", 1)[-1],
        "with_project": _norm_path(
            str(default_export_dir(project_path)), _ROOT),
    })

    # --- data_asset_registry ---------------------------------------------
    reg = DataAssetRegistry()
    reg.register_format(FormatSpec(
        format_id="xyzdata", extensions={"xyz"},
        resource_type="custom_type", status="indexed_custom",
        exporter=lambda asset, fid, out: (
            Path(out).write_text("spec-export", encoding="utf-8") or True),
    ))
    reg_classify = {}
    for name in ("table.csv", "x.xyz", "notes.txt"):
        reg_classify[name] = list(reg.classify_path(tree / name))
    scanned = reg.scan_directory(tree)
    reg_export_out = tree / "spec_out.bin"
    reg_export_ok = reg.export(
        ResourceItem(name="x.xyz", path=str(tree / "x.xyz"),
                     type="custom_type", format="xyz"),
        "xyzdata", reg_export_out)
    reg_export_table_out = tree / "conv_out.json"
    reg_export_table = reg.export(
        ResourceItem(name="table.csv", path=str(tree / "table.csv"),
                     type="tabular", format="csv"),
        "json", reg_export_table_out)
    try:
        reg.export(
            ResourceItem(name="x.xyz", path=str(tree / "x.xyz"),
                         type="custom_type", format="xyz"),
            "pdf", tree / "never.pdf")
        reg_export_err = None
    except Exception as e:
        reg_export_err = f"{e.__class__.__name__}: {e}"
    cases.append({
        "id": "registry",
        "classify": reg_classify,
        "scan_count": len(scanned),
        "scan_names": [r.name for r in scanned],
        "scan_types": [r.type for r in scanned],
        "spec_export_ok": reg_export_ok,
        "spec_export_text": reg_export_out.read_text(encoding="utf-8"),
        "conv_export_ok": reg_export_table,
        "conv_export_text": reg_export_table_out.read_text(encoding="utf-8"),
        "missing_exporter": reg_export_err,
    })

    OUT.write_text(
        json.dumps({"cases": cases}, ensure_ascii=False, indent=1),
        encoding="utf-8")
    print(f"frozen {OUT} ({len(cases)} cases)")


if __name__ == "__main__":
    main()
